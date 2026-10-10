"""Подготовка содержимого документа для провайдера: изображение, PDF или текст таблицы."""

import base64
import io
from dataclasses import dataclass

from openpyxl import load_workbook

MAX_TEXT_CHARS = 100_000
MAX_SHEETS, MAX_ROWS, MAX_COLS = 3, 500, 40


class InputError(Exception):
    """Содержимое не удалось подготовить (повреждённый файл и т.п.); повторять нет смысла."""


@dataclass(frozen=True)
class ProviderInput:
    kind: str  # "image" | "pdf" | "text"
    mime: str
    filename: str
    data: bytes | None = None
    text: str | None = None

    def data_url(self) -> str:
        assert self.data is not None
        return f"data:{self.mime};base64,{base64.b64encode(self.data).decode()}"


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1251", errors="replace")


def wrap_untrusted(text: str) -> str:
    """Текст документа — данные, не инструкции: явные границы + лимит размера."""
    body = text[:MAX_TEXT_CHARS]
    return f"<<<DOCUMENT DATA START>>>\n{body}\n<<<DOCUMENT DATA END>>>"


def _xlsx_text(data: bytes) -> str:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise InputError("xlsx_unreadable") from exc
    parts: list[str] = []
    try:
        for ws in wb.worksheets[:MAX_SHEETS]:
            parts.append(f"Sheet: {ws.title}")
            for n, row in enumerate(ws.iter_rows(values_only=True, max_col=MAX_COLS)):
                if n >= MAX_ROWS:
                    parts.append("... (rows truncated)")
                    break
                cells = ["" if c is None else str(c) for c in row]
                if any(cells):
                    parts.append("\t".join(cells))
    finally:
        wb.close()
    return "\n".join(parts)


def build_input(content_type: str, data: bytes, filename: str) -> ProviderInput:
    if content_type in ("image/png", "image/jpeg"):
        return ProviderInput("image", content_type, filename, data=data)
    if content_type == "application/pdf":
        return ProviderInput("pdf", content_type, filename, data=data)
    if content_type == "text/csv":
        return ProviderInput("text", "text/plain", filename, text=wrap_untrusted(_decode(data)))
    if content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return ProviderInput("text", "text/plain", filename, text=wrap_untrusted(_xlsx_text(data)))
    raise InputError("input_type_unsupported")
