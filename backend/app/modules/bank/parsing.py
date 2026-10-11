"""Разбор банковских выписок CSV/XLSX по явному профилю (сопоставлению колонок).

Форматы конкретных банков не зашиты: пользователь указывает, в какой колонке дата, сумма и т. д., формат даты и
числа. Всё нормализуется строго; строка, которую нельзя разобрать однозначно, — ошибка с кодом, а не «догадка».
"""

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel, model_validator

MAX_ROWS = 20_000
MAX_TEXT = 1000
DATE_FORMATS = ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y.%m.%d", "%m/%d/%Y")
_TIN_DIGITS = re.compile(r"[^0-9]")


class ImportProfile(BaseModel):
    format: str = "csv"  # csv | xlsx
    delimiter: str = ";"
    encoding: str = "utf-8"  # utf-8 | cp1251
    header_row: int = 1
    sheet: str | None = None
    date_column: str
    date_format: str = "%d.%m.%Y"
    amount_mode: str = (
        "signed"  # signed (плюс — поступление) | split (отдельные колонки прихода и расхода)
    )
    amount_column: str | None = None
    incoming_column: str | None = None
    outgoing_column: str | None = None
    decimal_separator: str = ","
    thousands_separator: str = " "
    currency_column: str | None = None
    description_column: str | None = None
    reference_column: str | None = None
    counterparty_name_column: str | None = None
    counterparty_tin_column: str | None = None

    @model_validator(mode="after")
    def _check(self) -> "ImportProfile":
        if self.format not in ("csv", "xlsx"):
            raise ValueError("profile_format_invalid")
        if self.delimiter not in (";", ",", "\t", "|"):
            raise ValueError("profile_delimiter_invalid")
        if self.encoding not in ("utf-8", "cp1251"):
            raise ValueError("profile_encoding_invalid")
        if not 1 <= self.header_row <= 100:
            raise ValueError("profile_header_row_invalid")
        if self.date_format not in DATE_FORMATS:
            raise ValueError("profile_date_format_invalid")
        if self.amount_mode not in ("signed", "split"):
            raise ValueError("profile_amount_mode_invalid")
        if self.amount_mode == "signed" and not self.amount_column:
            raise ValueError("profile_amount_column_required")
        if self.amount_mode == "split" and not (self.incoming_column and self.outgoing_column):
            raise ValueError("profile_split_columns_required")
        if self.decimal_separator not in (",", "."):
            raise ValueError("profile_decimal_invalid")
        if self.thousands_separator not in (" ", ",", ".", "", "'"):
            raise ValueError("profile_thousands_invalid")
        if self.thousands_separator == self.decimal_separator:
            raise ValueError("profile_separators_equal")
        return self


@dataclass
class ParsedRow:
    row_number: int
    txn_date: date
    amount: Decimal
    currency: str | None
    description: str | None
    reference: str | None
    counterparty_name: str | None
    counterparty_tin: str | None
    raw: dict[str, str]
    row_hash: str = ""


@dataclass
class RowError:
    row_number: int
    code: str
    detail: str = ""


@dataclass
class ParseResult:
    headers: list[str]
    rows: list[ParsedRow] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    total: int = 0


class ParseError(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


def _norm_header(v: Any) -> str:
    return str(v).strip() if v is not None else ""


def _read_csv(data: bytes, p: ImportProfile) -> list[list[str]]:
    try:
        text = data.decode("utf-8-sig" if p.encoding == "utf-8" else "cp1251")
    except UnicodeDecodeError as exc:
        raise ParseError(
            "encoding_mismatch", "Не удалось прочитать файл в выбранной кодировке."
        ) from exc
    reader = csv.reader(io.StringIO(text), delimiter=p.delimiter)
    rows: list[list[str]] = []
    try:
        for row in reader:
            rows.append(row)
            if len(rows) > MAX_ROWS + 200:
                raise ParseError("too_many_rows", f"В файле больше {MAX_ROWS} строк.")
    except csv.Error as exc:
        raise ParseError("csv_malformed", "Файл CSV разобрать не удалось.") from exc
    return rows


def _read_xlsx(data: bytes, p: ImportProfile) -> list[list[str]]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # повреждённая книга
        raise ParseError("xlsx_unreadable", "Не удалось прочитать книгу XLSX.") from exc
    try:
        ws: Any
        if p.sheet:
            if p.sheet not in wb.sheetnames:
                raise ParseError("sheet_not_found", "Лист с таким названием не найден.")
            ws = wb[p.sheet]
        else:
            ws = wb.worksheets[0]
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append([_xl_cell(v) for v in row])
            if len(rows) > MAX_ROWS + 200:
                raise ParseError("too_many_rows", f"В файле больше {MAX_ROWS} строк.")
        return rows
    finally:
        wb.close()


def _xl_cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        return format(Decimal(repr(v)), "f")  # без экспоненты; точное десятичное представление
    return str(v)


def read_table(data: bytes, p: ImportProfile) -> tuple[list[str], list[tuple[int, list[str]]]]:
    rows = _read_csv(data, p) if p.format == "csv" else _read_xlsx(data, p)
    if len(rows) < p.header_row:
        raise ParseError("header_not_found", "Строка заголовков не найдена.")
    headers = [_norm_header(h) for h in rows[p.header_row - 1]]
    body = [(i + p.header_row + 1, r) for i, r in enumerate(rows[p.header_row :])]
    body = [(n, r) for n, r in body if any(c.strip() for c in r)]  # пустые строки пропускаются
    if len(body) > MAX_ROWS:
        raise ParseError("too_many_rows", f"В файле больше {MAX_ROWS} строк.")
    return headers, body


def _parse_date(raw: str, p: ImportProfile) -> date:
    v = raw.strip()
    if p.format == "xlsx" and re.match(r"\d{4}-\d{2}-\d{2}", v):
        v = v[:10]  # ячейки-даты XLSX приходят в ISO независимо от отображаемого формата
        fmt = "%Y-%m-%d"
    else:
        fmt = p.date_format
    try:
        return datetime.strptime(v, fmt).date()
    except ValueError as exc:
        raise ValueError("date_invalid") from exc


def parse_amount(raw: str, p: ImportProfile) -> Decimal | None:
    """Сумма как напечатано в файле: пробелы (в т. ч. неразрывные) игнорируются, разделители — по профилю."""
    v = re.sub(r"\s+", "", raw.replace("\u00a0", " ").replace("\u202f", " "))
    if v == "":
        return None
    neg = False
    if v.startswith("(") and v.endswith(")"):
        neg, v = True, v[1:-1]
    if v.startswith("-"):
        neg, v = True, v[1:]
    elif v.startswith("+"):
        v = v[1:]
    if p.thousands_separator.strip():
        v = v.replace(p.thousands_separator, "")
    if p.decimal_separator == ",":
        v = v.replace(",", ".")
    if not re.fullmatch(r"[0-9]+(\.[0-9]+)?", v):
        raise ValueError("amount_invalid")
    try:
        d = Decimal(v)
    except InvalidOperation as exc:
        raise ValueError("amount_invalid") from exc
    return -d if neg else d


def _clean(v: str | None) -> str | None:
    if v is None:
        return None
    v = "".join(ch for ch in v if ch.isprintable() or ch in "\t").strip()
    return v[:MAX_TEXT] or None


def parse_statement(
    data: bytes, p: ImportProfile, *, currency: str, minor_units: int
) -> ParseResult:
    headers, body = read_table(data, p)
    index = {h: i for i, h in enumerate(headers) if h}
    needed = [p.date_column]
    if p.amount_mode == "signed":
        needed.append(p.amount_column or "")
    else:
        needed += [p.incoming_column or "", p.outgoing_column or ""]
    for opt in (
        p.currency_column,
        p.description_column,
        p.reference_column,
        p.counterparty_name_column,
        p.counterparty_tin_column,
    ):
        if opt:
            needed.append(opt)
    missing = [c for c in needed if c not in index]
    if missing:
        raise ParseError("columns_missing", "В файле нет колонок: " + ", ".join(missing))
    res = ParseResult(headers=[h for h in headers if h], total=len(body))

    def cell(row: list[str], col: str | None) -> str:
        if not col:
            return ""
        i = index[col]
        return row[i] if i < len(row) else ""

    seen: dict[tuple[str, ...], int] = {}
    for num, row in body:
        raw = {h: (row[i] if i < len(row) else "") for h, i in index.items()}
        try:
            d = _parse_date(cell(row, p.date_column), p)
            if not 2000 <= d.year <= 2100:
                raise ValueError("date_out_of_range")
            if p.amount_mode == "signed":
                amount = parse_amount(cell(row, p.amount_column), p)
                if amount is None:
                    raise ValueError("amount_missing")
            else:
                inc = parse_amount(cell(row, p.incoming_column), p)
                out = parse_amount(cell(row, p.outgoing_column), p)
                if (inc is not None and inc < 0) or (out is not None and out < 0):
                    raise ValueError("amount_negative_in_split")
                if (inc or 0) and (out or 0):
                    raise ValueError("amount_both_sides")
                if not (inc or 0) and not (out or 0):
                    raise ValueError("amount_missing")
                amount = inc if inc else -(out or Decimal(0))
            if amount == 0:
                raise ValueError("amount_zero")
            if -amount.normalize().as_tuple().exponent > minor_units:  # type: ignore[operator]
                raise ValueError("amount_too_many_decimals")
            cur = cell(row, p.currency_column).strip().upper() if p.currency_column else ""
            if cur and cur != currency:
                raise ValueError("currency_mismatch")
            tin_raw = cell(row, p.counterparty_tin_column)
            tin = _TIN_DIGITS.sub("", tin_raw) or None
        except ValueError as exc:
            res.errors.append(RowError(num, str(exc) if str(exc) else "row_invalid"))
            continue
        item = ParsedRow(
            row_number=num,
            txn_date=d,
            amount=amount,
            currency=cur or None,
            description=_clean(cell(row, p.description_column)),
            reference=_clean(cell(row, p.reference_column)),
            counterparty_name=_clean(cell(row, p.counterparty_name_column)),
            counterparty_tin=tin,
            raw=raw,
        )
        key = (
            d.isoformat(),
            format(amount, "f"),
            item.reference or "",
            item.description or "",
            item.counterparty_name or "",
            tin or "",
        )
        occ = seen.get(key, 0)
        seen[key] = occ + 1
        item.row_hash = hashlib.sha256("\x1f".join((*key, str(occ))).encode("utf-8")).hexdigest()
        res.rows.append(item)
    return res
